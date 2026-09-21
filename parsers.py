"""
Parsers for the known "result report" source formats.

Each parser reads a raw file and yields normalized dict records:
    {
        "candidate_name": str,
        "sa_id": str,
        "category": str,          # Citizenship | PERSAL | Criminal | Qualification | Matric
                                   # | SAQA-Matric | SAQA-Tertiary | UNISA | MUT
        "verification_type": str, # full display string, embeds qualification detail
        "breakdown_group": str | None,
        "supplier": str,          # display supplier for Consolidated Results
        "ver_date": str | None,
        "result_raw": str,
        "cost_excl": float,
        "cost_incl": float,
        "source_document": str,
        "notes": str,
    }
Raw rows that can't be priced under a known rule are returned separately as "skipped"
with a reason, so the app can show them instead of silently dropping data.
"""
import os
import re
import openpyxl
import pricing_rules as pr


def _clean(v):
    if v is None:
        return ""
    return str(v).strip()


def _name_from_parts(*parts):
    return " ".join(p for p in (_clean(x) for x in parts) if p).strip()


# ---------------------------------------------------------------------------
# MIE — Captured_Credentials.xlsx
# ---------------------------------------------------------------------------
def parse_mie_captured_credentials(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_row_idx = None
    for i, row in enumerate(rows[:5]):
        if row and _clean(row[0]).lower() in ("client",):
            header_row_idx = i
            break
    if header_row_idx is None:
        header_row_idx = 1  # fallback: row 2 (0-indexed 1) matches the known template

    headers = [_clean(h) for h in rows[header_row_idx]]
    col = {h: i for i, h in enumerate(headers)}
    fname = os.path.basename(path)

    records, skipped = [], []
    for row in rows[header_row_idx + 1:]:
        if row is None or all(v is None for v in row):
            continue
        get = lambda name, default=None: row[col[name]] if name in col and col[name] < len(row) else default

        sa_id = _clean(get("ID Number"))
        if not sa_id:
            continue
        name = _name_from_parts(get("Name "), get("Surname"))
        vtype_raw = _clean(get("Type"))
        supplier_raw = _clean(get("Supplier"))
        result_raw = _clean(get("Result"))
        ver_date = get("Result Date") or get("Date Captured")
        source_doc = f"{fname} (MIE)"
        tl = vtype_raw.lower()

        if "fraud check" in tl:
            skipped.append({"row": row, "reason": "Fraud Check — excluded per business rule (not priced)."})
            continue

        if "citizenship" in tl:
            excl, incl, breakdown, notes = pr.price_citizenship_mie()
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Citizenship",
                verification_type="Citizenship - Citizenship - South Africa",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "persal" in tl:
            excl, incl, breakdown, notes = pr.price_persal_mie()
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="PERSAL",
                verification_type="PERSAL - PERSAL Public Sector Employment Status",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "fingerprintzone" in tl:
            excl, incl, breakdown, notes = pr.price_criminal()
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Criminal",
                verification_type="Criminal - FingerprintZone (Standard) - South Africa",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "umalusi" in tl and "national secondary" in tl:
            excl, incl, breakdown, notes = pr.price_matric("National Secondary")
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Matric",
                verification_type=f"Matric - UMALUSI - National Secondary - South Africa ({_clean(get('Attr1')) or 'Matric'})",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "umalusi" in tl and "matric symbol match" in tl:
            excl, incl, breakdown, notes = pr.price_matric("Symbol Match")
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Matric",
                verification_type=f"Matric - UMALUSI - Matric Symbol Match - South Africa ({_clean(get('Attr1')) or 'Matric'})",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "umalusi" in tl and "tertiary symbol match" in tl:
            excl, incl, _, _ = pr.price_matric("Symbol Match")
            notes = (
                "ASSUMPTION: UMALUSI - Tertiary Symbol Match treated the same as Matric Symbol Match "
                f"(UMALUSI base R{pr.MATRIC_BASE:.2f} + Verification/Symbol Match Fee R{pr.MATRIC_SYMBOL_MATCH_FEE:.2f} "
                f"= R{excl:.2f}) - not explicitly confirmed. | Incl. 15% VAT: R{excl:.2f} x 1.15 = R{incl:.2f}"
            )
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Matric",
                verification_type=f"Matric - UMALUSI - Tertiary Symbol Match - South Africa ({_clean(get('Attr1')) or ''})".strip(),
                breakdown_group="Symbol Match (Tertiary variant - assumed same rate)", supplier="MIE",
                ver_date=ver_date, result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif tl.startswith("qualification general"):
            country = vtype_raw.split("-", 1)[-1].strip() if "-" in vtype_raw else ""
            qual = _clean(get("Attr1"))
            excl, incl, breakdown, notes = pr.price_qualification("National Qualifications Register", "Examinations Council of Lesotho")
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Qualification",
                verification_type=f"Qualification - {vtype_raw} ({qual})" if qual else f"Qualification - {vtype_raw}",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        elif "national qualifications register" in tl or "national tertiary" in tl or "tertiary course" in tl:
            kind = ("National Qualifications Register" if "national qualifications register" in tl
                    else "National Tertiary" if "national tertiary" in tl else "Tertiary Course")
            qual = _clean(get("Attr1"))
            excl, incl, breakdown, notes = pr.price_qualification(kind, supplier_raw)
            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="Qualification",
                verification_type=f"Qualification - {kind} - South Africa ({qual})" if qual else f"Qualification - {kind} - South Africa",
                breakdown_group=breakdown, supplier="MIE", ver_date=ver_date,
                result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=source_doc, notes=notes,
            ))
        else:
            skipped.append({"row": row, "reason": f"Unrecognized MIE verification type: '{vtype_raw}' — no pricing rule."})

    return records, skipped


# ---------------------------------------------------------------------------
# SAQA / NLRD PDF report
# ---------------------------------------------------------------------------
def parse_saqa_pdf(path, mode="bulk"):
    """
    mode: 'bulk' (R58 flat) or 'individual' (R82 flat) — the user states which,
    per the business rule that governs SAQA submission-type pricing.
    Requires the pdfplumber package (falls back to a helpful error if missing).
    """
    try:
        import pdfplumber
    except ImportError as e:
        raise RuntimeError("pdfplumber is required to parse SAQA PDF reports. Install it with: python -m pip install pdfplumber") from e

    fname = os.path.basename(path)
    records, skipped = [], []
    id_pattern = re.compile(r"^\d{13}$")

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                if not table:
                    continue
                header = [_clean(c).lower() for c in table[0]]
                if not any("national id" in h for h in header):
                    continue
                idx = {}
                for i, h in enumerate(header):
                    if "company" in h:
                        idx["company"] = i
                    elif h == "name":
                        idx["name"] = i
                    elif "national id" in h:
                        idx["id"] = i
                    elif "qualification" in h:
                        idx["qual"] = i
                    elif "institution" in h:
                        idx["institution"] = i
                    elif h == "year":
                        idx["year"] = i
                    elif "found" in h:
                        idx["found"] = i
                    elif "feedback" in h or "comment" in h:
                        idx["comment"] = i
                    elif "batch" in h:
                        idx["batch"] = i

                for row in table[1:]:
                    if not row or "id" not in idx:
                        continue
                    sa_id = _clean(row[idx["id"]])
                    if not id_pattern.match(sa_id):
                        continue
                    name = _clean(row[idx.get("name", -1)]) if idx.get("name") is not None else ""
                    qual = _clean(row[idx.get("qual", -1)]) if idx.get("qual") is not None else ""
                    found = _clean(row[idx.get("found", -1)]) if idx.get("found") is not None else ""
                    batch = _clean(row[idx.get("batch", -1)]) if idx.get("batch") is not None else ""

                    is_matric = qual.strip().lower() in ("senior certificate", "matric", "national senior certificate")
                    category = "SAQA-Matric" if is_matric else "SAQA-Tertiary"
                    excl, incl, breakdown, notes = pr.price_saqa(mode)
                    vtype = f"{category} - {category} ({qual})" if qual else f"{category} - {category}"
                    source_doc = f"{fname} (SAQA/NLRD{', batch ' + batch if batch else ''})"

                    records.append(dict(
                        candidate_name=name, sa_id=sa_id, category=category,
                        verification_type=vtype, breakdown_group=breakdown, supplier="SAQA",
                        ver_date=None, result_raw=found or "Yes", cost_excl=excl, cost_incl=incl,
                        source_document=source_doc, notes=notes,
                    ))
    return records, skipped


# ---------------------------------------------------------------------------
# UNISA — R4C upload / bulk order files
# ---------------------------------------------------------------------------
def parse_unisa_r4c(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    fname = os.path.basename(path)
    records, skipped = [], []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [_clean(h) for h in rows[0]]
        col = {h: i for i, h in enumerate(headers)}
        if "ID Number" not in col and "Id Number" not in col:
            continue
        id_col = col.get("ID Number", col.get("Id Number"))

        for row in rows[1:]:
            if row is None or all(v is None for v in row):
                continue
            get = lambda name, default=None: row[col[name]] if name in col and col[name] < len(row) else default
            sa_id = _clean(row[id_col]) if id_col is not None else ""
            if not sa_id:
                continue
            name = _name_from_parts(get("First Name(s)"), get("Surname"))
            qual = _clean(get("Qualification Name"))
            result_raw = _clean(get("Result"))
            ver_date = get("Verified At")
            comments = _clean(get("Comments"))

            excl, incl, detail_note = pr.price_direct_free("UNISA", f"via R4C upload ({fname})")
            vtype = f"UNISA - UNISA ({qual})" if qual else "UNISA - UNISA"
            notes = "Direct UNISA verification via R4C upload - free per pricing rules."
            if comments:
                notes += f" Comments: {comments}."

            records.append(dict(
                candidate_name=name, sa_id=sa_id, category="UNISA",
                verification_type=vtype, breakdown_group=None, supplier="Direct - UNISA",
                ver_date=ver_date, result_raw=result_raw, cost_excl=excl, cost_incl=incl,
                source_document=f"{fname} (R4C upload)", notes=notes,
            ))
    return records, skipped


# ---------------------------------------------------------------------------
# MUT — Candidates file
# ---------------------------------------------------------------------------
def parse_mut_candidates(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [_clean(h) for h in rows[0]]
    col = {h: i for i, h in enumerate(headers)}
    fname = os.path.basename(path)
    records, skipped = [], []

    id_key = next((h for h in headers if "id number" in h.lower()), None)
    complete_key = next((h for h in headers if "complete" in h.lower()), None)
    comment_key = next((h for h in headers if "comment" in h.lower() or "graduation" in h.lower()), None)

    for row in rows[1:]:
        if row is None or all(v is None for v in row):
            continue
        get = lambda name, default=None: row[col[name]] if name in col and col[name] < len(row) else default
        sa_id = _clean(row[col[id_key]]) if id_key else ""
        if not sa_id:
            continue
        name = _name_from_parts(get("NAMES"), get("SURNAME"))
        qual = _clean(get("QUALIFICATION NAME"))
        complete = _clean(row[col[complete_key]]) if complete_key else ""
        comment_raw = row[col[comment_key]] if comment_key else None
        comment = _clean(comment_raw)

        excl, incl, _ = pr.price_direct_free("MUT", "")
        vtype = f"MUT - MUT ({qual})" if qual else "MUT - MUT"
        notes = "Direct MUT verification - free per pricing rules."
        if complete:
            notes += f" Completion status: {complete}."
        if comment:
            notes += f" Comment: {comment}"

        records.append(dict(
            candidate_name=name, sa_id=sa_id, category="MUT",
            verification_type=vtype, breakdown_group=None, supplier="Direct - MUT",
            ver_date=comment_raw, result_raw=complete, cost_excl=excl, cost_incl=incl,
            source_document=f"{fname} (Direct MUT)", notes=notes,
        ))
    return records, skipped


def detect_and_parse(path, saqa_mode="bulk"):
    """Best-effort format auto-detection by filename + header shape."""
    lower = os.path.basename(path).lower()
    if lower.endswith(".pdf"):
        return parse_saqa_pdf(path, mode=saqa_mode), "SAQA/NLRD PDF"

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    first_row = next(ws.iter_rows(max_row=1, values_only=True), ())
    second_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), ())
    headers = [_clean(h).lower() for h in (first_row or ())]
    headers2 = [_clean(h).lower() for h in (second_row or ())]
    wb.close()

    if "client" in headers2 and "id number" in headers2:
        return parse_mie_captured_credentials(path), "MIE Captured Credentials"
    if any("id number" in h for h in headers) and any("verification type" in h for h in headers):
        return parse_unisa_r4c(path), "UNISA R4C Upload"
    if any("id number" in h for h in headers) and any("qualification name" in h for h in headers):
        return parse_mut_candidates(path), "MUT Candidates"
    raise ValueError(f"Could not auto-detect format for {os.path.basename(path)}. "
                      "Recognized formats: MIE Captured Credentials (.xlsx), SAQA/NLRD report (.pdf), "
                      "UNISA R4C upload (.xlsx), MUT Candidates (.xlsx).")
