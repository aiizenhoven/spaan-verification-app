"""
The ledger engine: loads the existing "Consolidated Results" sheet, appends newly
parsed records with ingestion-dedup + within-data duplicate flagging, and applies
the assumed Citizenship/PERSAL fallback rule for candidates who still lack one.
"""
import openpyxl
import pricing_rules as pr

CONSOLIDATED_HEADERS = [
    "Candidate Name", "South African ID Number", "Verification Type", "Category",
    "Breakdown Group", "Verification Supplier", "Verification Date", "Verification Result",
    "Verification Status", "Cost (Excl. VAT - as sourced)", "Cost (Incl. VAT)",
    "Duplicate?", "Source Document", "Notes",
]


def load_ledger(master_path):
    """Returns list of dict rows from the existing Consolidated Results sheet."""
    wb = openpyxl.load_workbook(master_path, data_only=True)
    ws = wb["Consolidated Results"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    ledger = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["Candidate Name"]] is None:
            continue
        ledger.append({h: row[idx[h]] for h in CONSOLIDATED_HEADERS if h in idx})
    return ledger


def _ingestion_key(rec):
    return (
        str(rec["sa_id"]).strip(),
        rec["verification_type"].strip(),
        str(rec.get("ver_date") or "").strip(),
        str(rec.get("result_raw") or "").strip(),
    )


def _existing_ingestion_keys(ledger):
    keys = set()
    for row in ledger:
        keys.add((
            str(row.get("South African ID Number") or "").strip(),
            str(row.get("Verification Type") or "").strip(),
            str(row.get("Verification Date") or "").strip(),
            str(row.get("Verification Result") or "").strip(),
        ))
    return keys


def _occurrence_key(rec_or_row, from_ledger=False):
    if from_ledger:
        return (str(rec_or_row.get("South African ID Number") or "").strip(),
                str(rec_or_row.get("Verification Type") or "").strip())
    return (str(rec_or_row["sa_id"]).strip(), rec_or_row["verification_type"].strip())


def merge_batch(ledger, new_records):
    """
    Applies ingestion-dedup and duplicate flagging. Returns:
        (added_rows, skipped_as_already_ingested)
    added_rows are dicts matching CONSOLIDATED_HEADERS, ready to append.
    Does NOT mutate `ledger`; caller should extend it with added_rows afterwards.
    """
    existing_ingestion_keys = _existing_ingestion_keys(ledger)
    occurrence_counts = {}
    for row in ledger:
        k = _occurrence_key(row, from_ledger=True)
        occurrence_counts[k] = occurrence_counts.get(k, 0) + 1

    added, already_ingested = [], []
    seen_this_batch = set()

    for rec in new_records:
        ikey = _ingestion_key(rec)
        if ikey in existing_ingestion_keys or ikey in seen_this_batch:
            already_ingested.append(rec)
            continue
        seen_this_batch.add(ikey)

        okey = _occurrence_key(rec)
        occurrence_counts[okey] = occurrence_counts.get(okey, 0) + 1
        is_duplicate = occurrence_counts[okey] > 1

        added.append({
            "Candidate Name": rec["candidate_name"],
            "South African ID Number": rec["sa_id"],
            "Verification Type": rec["verification_type"],
            "Category": rec["category"],
            "Breakdown Group": rec.get("breakdown_group"),
            "Verification Supplier": rec["supplier"],
            "Verification Date": rec.get("ver_date"),
            "Verification Result": rec.get("result_raw"),
            "Verification Status": "Located",
            "Cost (Excl. VAT - as sourced)": rec["cost_excl"],
            "Cost (Incl. VAT)": rec["cost_incl"],
            "Duplicate?": "Yes" if is_duplicate else "No",
            "Source Document": rec["source_document"],
            "Notes": rec["notes"],
        })

    return added, already_ingested


def apply_assumed_fallbacks(ledger):
    """
    For every candidate in the ledger with no Citizenship row, add an assumed
    Datanamics API Citizenship row. Same for PERSAL -> API. Idempotent: only
    adds if genuinely absent. Returns the list of new fallback rows.
    """
    by_id = {}
    for row in ledger:
        sa_id = str(row.get("South African ID Number") or "").strip()
        if not sa_id:
            continue
        by_id.setdefault(sa_id, {"name": row.get("Candidate Name"), "categories": set()})
        by_id[sa_id]["categories"].add(row.get("Category"))

    fallback_rows = []
    for sa_id, info in by_id.items():
        if "Citizenship" not in info["categories"]:
            excl, incl, breakdown, notes = pr.price_citizenship_assumed()
            fallback_rows.append({
                "Candidate Name": info["name"], "South African ID Number": sa_id,
                "Verification Type": "Citizenship - Citizenship (Datanamics API - Assumed)",
                "Category": "Citizenship", "Breakdown Group": breakdown,
                "Verification Supplier": "Datanamics API", "Verification Date": None,
                "Verification Result": "Assumed", "Verification Status": "Assumed (no source record)",
                "Cost (Excl. VAT - as sourced)": excl, "Cost (Incl. VAT)": incl,
                "Duplicate?": "No",
                "Source Document": "Assumed - no MIE Citizenship record found (business rule)",
                "Notes": notes,
            })
            info["categories"].add("Citizenship")
        if "PERSAL" not in info["categories"]:
            excl, incl, breakdown, notes = pr.price_persal_assumed()
            fallback_rows.append({
                "Candidate Name": info["name"], "South African ID Number": sa_id,
                "Verification Type": "PERSAL - PERSAL (API - Assumed)",
                "Category": "PERSAL", "Breakdown Group": breakdown,
                "Verification Supplier": "API", "Verification Date": None,
                "Verification Result": "Assumed", "Verification Status": "Assumed (no source record)",
                "Cost (Excl. VAT - as sourced)": excl, "Cost (Incl. VAT)": incl,
                "Duplicate?": "No",
                "Source Document": "Assumed - no MIE PERSAL record found (business rule)",
                "Notes": notes,
            })
            info["categories"].add("PERSAL")

    return fallback_rows
